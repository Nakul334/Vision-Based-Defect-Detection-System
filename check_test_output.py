import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\test_output.docx')
print("Sections count:", len(doc.sections))
for i, sec in enumerate(doc.sections):
    print(f"Section {i}: start_type={sec.start_type}, footer count={len(sec.footer.paragraphs)}")
    for fp in sec.footer.paragraphs:
        print(f"  Footer text: {repr(fp.text)}")
