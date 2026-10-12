import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')
print("Initial paragraphs:", len(doc.paragraphs))
print("Initial tables:", len(doc.tables))
print("Initial sections:", len(doc.sections))
