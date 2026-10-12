import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')
print('Num paragraphs:', len(doc.paragraphs))
print('Num tables:', len(doc.tables))
print('Num sections:', len(doc.sections))

with open(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\template_dump.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'P{i} [{p.style.name}]: {p.text}\n')
    for t_idx, t in enumerate(doc.tables):
        f.write(f'\n--- TABLE {t_idx} ({len(t.rows)}x{len(t.columns)}) ---\n')
        for r_idx, row in enumerate(t.rows):
            row_txt = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
            f.write(f'Row {r_idx}: ' + ' | '.join(row_txt) + '\n')

print('Dumped template info successfully.')
