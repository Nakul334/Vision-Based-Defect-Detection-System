import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')

print("Available Sections:", len(doc.sections))
for i, sec in enumerate(doc.sections):
    print(f"Section {i}: top={sec.top_margin.inches}in, bottom={sec.bottom_margin.inches}in, left={sec.left_margin.inches}in, right={sec.right_margin.inches}in")
    print(f"  header p count: {len(sec.header.paragraphs)}, footer p count: {len(sec.footer.paragraphs)}")
    for hp in sec.header.paragraphs:
        if hp.text.strip():
            print(f"    Header: {hp.text}")
    for fp in sec.footer.paragraphs:
        if fp.text.strip():
            print(f"    Footer: {fp.text}")

print("\nStyles check:")
for s_name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3']:
    try:
        st = doc.styles[s_name]
        print(f"  Style '{s_name}': font={st.font.name}, size={st.font.size}")
    except Exception as e:
        print(f"  Style '{s_name}' error: {e}")
