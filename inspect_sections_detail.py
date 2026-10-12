import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')

for s_idx, sec in enumerate(doc.sections):
    print(f"--- Section {s_idx} ---")
    print("Page width:", sec.page_width.inches, "height:", sec.page_height.inches)
    print("Margins: top=", sec.top_margin.inches, "bottom=", sec.bottom_margin.inches, "left=", sec.left_margin.inches, "right=", sec.right_margin.inches)
    print("Different first page header/footer:", sec.different_first_page_header_footer)
    for p in sec.header.paragraphs:
        print("  Header text:", repr(p.text))
    for p in sec.footer.paragraphs:
        print("  Footer text:", repr(p.text))
