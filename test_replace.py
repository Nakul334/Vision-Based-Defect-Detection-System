import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')

# Check how easy it is to replace text in P0
print("Original P0:", doc.paragraphs[0].text)
doc.paragraphs[0].text = "AN AUTOMATED ROBOTIC VISION SYSTEM FOR PRECISION SURFACE DEFECT INSPECTION IN INDUSTRIAL PACKAGING USING YOLOV8-OBB"
print("New P0:", doc.paragraphs[0].text)

# Check P2
print("P2 xml length:", len(doc.paragraphs[2]._p.xml))
