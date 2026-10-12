import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')

children = list(doc._body._element)
print("Before children count:", len(children))

# Save the sectPr from the end if needed
last_sectPr = children[-1]

# Delete children from index 82 to end-1
for c in children[82:-1]:
    doc._body._element.remove(c)

print("After children count:", len(list(doc._body._element)))

# Try adding a heading and paragraph
p = doc.add_paragraph("Chapter 1: Introduction", style='Heading 1')
p2 = doc.add_paragraph("This is the full introduction text test.")

doc.save(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\test_output.docx')
print("Successfully saved test_output.docx")
