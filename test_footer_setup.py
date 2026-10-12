import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_page_number_footer(sec, fmt="decimal", start=1):
    sec.different_first_page_header_footer = False
    footer = sec.footer
    footer.is_linked_to_previous = False
    
    # Configure pgNumType on section
    sectPr = sec._sectPr
    # remove existing pgNumType if any
    for child in sectPr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgNumType'):
        sectPr.remove(child)
    
    pgNum = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}" w:start="{start}"/>')
    sectPr.append(pgNum)
    
    # Clear footer paragraphs
    for p in footer.paragraphs:
        p._p.getparent().remove(p._p)
    
    # Add new paragraph with centered PAGE field
    p = footer.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p_pr = p._p.get_or_add_pPr()
    
    # Add page number run
    fld_xml = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'  <w:pPr><w:jc w:val="center"/></w:pPr>'
        f'  <w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r>'
        f'  <w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        f'  <w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr><w:fldChar w:fldCharType="separate"/></w:r>'
        f'  <w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr><w:t>1</w:t></w:r>'
        f'  <w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="20"/></w:rPr><w:fldChar w:fldCharType="end"/></w:r>'
        f'</w:p>'
    )
    p._p.getparent().replace(p._p, fld_xml)

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\test_output.docx')
set_page_number_footer(doc.sections[2], fmt="decimal", start=1)
doc.save(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\test_output2.docx')
print("Successfully tested footer setup!")
