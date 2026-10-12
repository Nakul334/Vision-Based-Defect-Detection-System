import docx

doc = docx.Document(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\Major_project Report_Templet.docx')

with open(r'C:\Users\nakul\Downloads\Major_P_CJ_Rice\Rice_Case_01.yolov8\all_paragraphs.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        f.write(f"P{i:03d} [{p.style.name}] (runs={len(p.runs)}): {repr(p.text)}\n")

print(f"Dumped {len(doc.paragraphs)} paragraphs to all_paragraphs.txt")
